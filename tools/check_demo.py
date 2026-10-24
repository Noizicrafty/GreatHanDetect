"""Read-only acceptance smoke check against a running demo; no events are saved."""
import argparse
import base64
import hashlib
import io
import json
import urllib.request
import zipfile

CHECKPOINT_SHA256 = '5b9df5d26785a3848643b7603036c2219dea4aedd6a19373dadaaa56182e519d'


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--url', default='http://127.0.0.1:8090')
    p.add_argument('--health-only', action='store_true')
    p.add_argument('--tpu', action='store_true', help='Also require actual LinQ inference')
    a = p.parse_args()
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    checks = []

    def request(path, body=None):
        req = urllib.request.Request(a.url.rstrip('/') + path,
            data=json.dumps(body).encode() if body is not None else None,
            headers={'Content-Type': 'application/json'})
        with opener.open(req, timeout=90) as r:
            assert r.status == 200, (path, r.status)
            return r.read()

    before = json.loads(request('/api/demo/state'))
    assert before['demo'] and before['model_ready'] and not before['training_enabled'], before
    checks.append('HTTP ready, checkpoint present, training disabled')
    if a.health_only:
        print(json.dumps({'healthy': True}))
        return
    for page in ('03_cameras', '01_monitoring', '02_analytics', '04_notifications', '05_report'):
        assert b'demo.js' in request('/' + page + '.html')
        checks.append(page)
    for path in ('webcam.mjs', 'webcam-face-worker.mjs', 'assets/mediapipe/face_landmarker.task',
                 'assets/mediapipe/wasm/vision_wasm_internal.wasm'):
        assert request('/' + path)
    checks.append('Webcam and face assets served locally')
    from pypdf import PdfReader
    pdf = request('/api/demo/report.pdf?download=1')
    assert pdf.startswith(b'%PDF')
    text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert 'Отчёт по инцидентам' in text, text[:200]
    checks.append('PDF contains Cyrillic text')
    from openpyxl import load_workbook
    assert load_workbook(io.BytesIO(request('/api/demo/export.xlsx'))).active.max_row >= 1
    checks.append('XLSX opens')
    with zipfile.ZipFile(io.BytesIO(request('/api/demo/dataset.zip'))) as z:
        assert z.testzip() is None
    checks.append('Dataset ZIP opens')
    raw = request('/assets/photos/trained_sample.jpg')
    positive = {'image': base64.b64encode(raw).decode()}
    prediction = json.loads(request('/api/demo/webcam', positive))
    assert prediction['prediction']['checkpoint_sha256'] == CHECKPOINT_SHA256
    assert prediction['prediction']['epochs'] == 26 and not prediction['saved']
    assert any(d['class_name'] == 'cigarette' for d in prediction['prediction']['detections'])
    checks.append('Actual 26-epoch CPU model detects cigarette on supplied sample')
    from PIL import Image
    blank = io.BytesIO()
    Image.new('RGB', (320, 240), (10, 15, 20)).save(blank, format='JPEG')
    negative = json.loads(request('/api/demo/webcam', {'image': base64.b64encode(blank.getvalue()).decode()}))
    assert negative['prediction']['detections'] == [] and not negative['saved']
    checks.append('Blank image produces no detections')
    if a.tpu:
        health = json.loads(request('/api/demo/linq-health'))
        assert health['ready'] and health['backend'] == 'linq_amqp'
        actual = json.loads(request('/api/demo/webcam/linq', positive))
        assert actual['prediction']['backend'] == 'linq_amqp'
        assert actual['prediction']['checkpoint_sha256'] == CHECKPOINT_SHA256
        assert actual['prediction']['program_sha256'] == 'f2920f42e078ca3e3e7df2bc6755c96a4fa823cf42873816c0b236fcaa048b5c'
        assert actual['prediction']['timing']['request_id']
        checks.append('Actual LinQ response, program/checkpoint hashes and request ID')
    after = json.loads(request('/api/demo/state'))
    assert before['events'] == after['events'], 'Smoke check changed journal'
    checks.append('Journal unchanged')
    print(json.dumps({'passed': True, 'checks': checks, 'model_sha256': CHECKPOINT_SHA256}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
