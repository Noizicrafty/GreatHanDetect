"""Integration checks for the presentation demo; uses an isolated temporary journal.

Run with the document QA runtime (Pillow, reportlab, openpyxl, PyMuPDF).
The inference-only tests require .venv-model and the existing checkpoint.
"""
import argparse,base64,concurrent.futures,importlib.util,io,json,sys,tempfile,threading,time,unittest,urllib.request,urllib.error,zipfile
import os
from pathlib import Path
from types import SimpleNamespace
from http.server import ThreadingHTTPServer
from html.parser import HTMLParser

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('presentation_demo',ROOT/'tools/run_presentation_demo.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
OUT=ROOT/'results/presentation_demo_qa';OUT.mkdir(exist_ok=True)

class DemoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.args=SimpleNamespace(root=str(ROOT/'demo/ff'),data=cls.temp.name,model_python=os.environ.get('DEMO_MODEL_PYTHON',str(ROOT/'.venv-model/Scripts/python.exe')),checkpoint=os.environ.get('DEMO_CHECKPOINT',str(ROOT/'results/training_roboflow_v4/smoking/weights/last.pt')))
        cls.app=module.Demo(cls.args);handler=module.make_handler(cls.app);handler.log_message=lambda *_:None
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),handler);cls.base=f'http://127.0.0.1:{cls.server.server_port}'
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close();cls.thread.join();cls.app.close_webcam();cls.temp.cleanup()
    def setUp(self):
        with self.app.connect() as db:db.execute('DELETE FROM events')
        self.app.reset()
    def request(self,path,data=None,headers=None,raw=None):
        payload=raw if raw is not None else (json.dumps(data).encode() if data is not None else None)
        req=urllib.request.Request(self.base+path,data=payload,headers=headers or {})
        try:
            with urllib.request.urlopen(req,timeout=100) as r:return r.status,r.headers,r.read()
        except urllib.error.HTTPError as e:return e.code,e.headers,e.read()
    def get_state(self):return json.loads(self.request('/api/demo/state')[2])
    def test_01_root_and_five_routes(self):
        for name in ['03_cameras.html','01_monitoring.html','02_analytics.html','04_notifications.html','05_report.html']:
            with self.subTest(page=name):
                code,headers,body=self.request('/'+name);self.assertEqual(code,200);self.assertIn(b'demo.js',body);self.assertIn('text/html',headers['Content-Type'])
        self.assertEqual(self.request('/')[0],200)
    def test_02_all_referenced_assets_exist(self):
        class Parser(HTMLParser):
            def __init__(self):super().__init__();self.paths=set()
            def handle_starttag(self,tag,attrs):
                a=dict(attrs)
                for k in ('src','href'):
                    if a.get(k) and not a[k].startswith(('http','#','data:')):self.paths.add(a[k].split('#')[0])
        for file in (ROOT/'demo/ff').glob('*.html'):
            parser=Parser();parser.feed(file.read_text(encoding='utf-8'))
            for path in parser.paths:
                with self.subTest(page=file.name,asset=path):self.assertEqual(self.request('/'+path.lstrip('/'))[0],200)
    def test_03_state_is_explicit_demo_and_five_events(self):
        s=self.get_state();self.assertTrue(s['demo']);self.assertEqual(len(s['events']),5);self.assertEqual(s['pending'],2);self.assertFalse(s['training_enabled']);self.assertTrue(all(e['source']=='presentation_demo' for e in s['events']))
    def test_04_confirm_is_persistent_and_idempotent(self):
        for _ in range(2):self.assertEqual(self.request('/api/demo/review',{'id':'EVT-2048','status':'confirmed'})[0],200)
        second=module.Demo(self.args);e=next(e for e in second.events() if e['id']=='EVT-2048');self.assertEqual(e['status'],'confirmed');self.assertEqual(len(second.events()),5);self.assertEqual(self.get_state()['pending'],1)
    def test_05_reject_updates_journal(self):
        self.assertEqual(self.request('/api/demo/review',{'id':'EVT-2046','status':'rejected'})[0],200);self.assertEqual(next(e for e in self.get_state()['events'] if e['id']=='EVT-2046')['status'],'rejected')
    def test_06_unknown_and_invalid_review_do_not_mutate(self):
        self.assertEqual(self.request('/api/demo/review',{'id':'missing','status':'confirmed'})[0],404)
        self.assertEqual(self.request('/api/demo/review',{'id':'EVT-2048','status':'made-up'})[0],400);self.assertEqual(self.get_state()['pending'],2)
    def test_07_reset_restores_fixtures(self):
        self.request('/api/demo/review',{'id':'EVT-2048','status':'rejected'});self.assertEqual(self.request('/api/demo/reset',{})[0],200);self.assertEqual(self.get_state()['pending'],2)
    def test_08_invalid_json(self):self.assertEqual(self.request('/api/demo/review',raw=b'{broken')[0],400)
    def test_09_bad_origin_and_host(self):
        self.assertEqual(self.request('/api/demo/reset',{},headers={'Origin':'https://example.invalid'})[0],403)
        self.assertEqual(self.request('/api/demo/state',headers={'Host':'example.invalid'})[0],403)
    def test_10_traversal_and_unknown_routes(self):
        for p in ['/../tools/run_presentation_demo.py','/%2e%2e/tools/run_presentation_demo.py','/demo-uploads/../demo.sqlite3','/api/no-such-route']:
            with self.subTest(path=p):self.assertEqual(self.request(p)[0],404)
    def test_11_oversized_request(self):self.assertEqual(self.request('/api/demo/review',raw=b'{}',headers={'Content-Length':str(16*1024*1024)})[0],413)
    def test_12_pdf_download_headers_and_binary(self):
        code,h,b=self.request('/api/demo/report.pdf?download=1');self.assertEqual(code,200);self.assertEqual(h['Content-Type'],'application/pdf');self.assertEqual(h['Content-Disposition'],'attachment; filename="LinQ_Safety_report.pdf"');self.assertEqual(int(h['Content-Length']),len(b));self.assertTrue(b.startswith(b'%PDF-'));self.assertTrue(b.rstrip().endswith(b'%%EOF'));(OUT/'download.pdf').write_bytes(b)
    def test_13_pdf_preview_matches_download_content(self):
        import pymupdf
        _,h,preview=self.request('/api/demo/report.pdf');_,_,download=self.request('/api/demo/report.pdf?download=1');self.assertNotIn('Content-Disposition',h)
        a=pymupdf.open(stream=preview,filetype='pdf');b=pymupdf.open(stream=download,filetype='pdf');self.assertEqual([p.get_text() for p in a],[p.get_text() for p in b]);self.assertEqual(len(a),1)
    def test_14_pdf_cyrillic_and_counts_follow_review(self):
        import pymupdf
        self.request('/api/demo/review',{'id':'EVT-2048','status':'confirmed'});self.request('/api/demo/review',{'id':'EVT-2046','status':'rejected'})
        data=self.request('/api/demo/report.pdf')[2];doc=pymupdf.open(stream=data,filetype='pdf');text='\n'.join(p.get_text() for p in doc)
        self.assertIn('Отчёт по инцидентам',text);self.assertIn('5\n2\n1\n0',text);self.assertIn('Отклонено',text);self.assertNotIn('\ufffd',text)
        doc[0].get_pixmap(matrix=pymupdf.Matrix(1.4,1.4)).save(str(OUT/'report_reviewed.png'))
    def test_15_pdf_multiple_pages_and_no_out_of_bounds_text(self):
        import pymupdf
        with self.app.connect() as db:
            for i in range(100):
                e={**module.SEED[i%5],'id':f'PAGINATION-{i}','date':'23.10.2026','source':'presentation_demo','title':'Проверка переноса длинного названия события и строк таблицы '+str(i)}
                db.execute('INSERT INTO events VALUES(?,?)',(e['id'],json.dumps(e,ensure_ascii=False)))
        data=self.request('/api/demo/report.pdf')[2];(OUT/'pagination.pdf').write_bytes(data);doc=pymupdf.open(stream=data,filetype='pdf');self.assertGreater(len(doc),2)
        for i,page in enumerate(doc):
            for word in page.get_text('words'):self.assertGreaterEqual(word[0],0);self.assertLessEqual(word[2],page.rect.width+.5);self.assertLessEqual(word[3],page.rect.height+.5)
            page.get_pixmap(matrix=pymupdf.Matrix(.8,.8)).save(str(OUT/f'pagination_{i+1}.png'))
        self.assertIn('99','\n'.join(p.get_text() for p in doc))
    def test_16_xlsx_opens_and_matches_journal(self):
        from openpyxl import load_workbook
        code,h,b=self.request('/api/demo/export.xlsx');self.assertEqual(code,200);self.assertTrue(b.startswith(b'PK'));wb=load_workbook(io.BytesIO(b),data_only=False);ws=wb['События'];self.assertEqual(ws.max_row,6);self.assertEqual(ws['E2'].number_format,'0%');self.assertEqual(ws.freeze_panes,'A2');self.assertTrue(all(row[6]=='DEMO презентации' for row in ws.iter_rows(min_row=2,values_only=True)));(OUT/'events.xlsx').write_bytes(b)
    def test_17_xlsx_category_and_search_filters(self):
        from openpyxl import load_workbook
        for query,count in [('?category=helmet',2),('?category=smoking',1),('?category=helmet&q=CAM-04',0),('?q=does-not-exist',0)]:
            with self.subTest(query=query):self.assertEqual(load_workbook(io.BytesIO(self.request('/api/demo/export.xlsx'+query)[2])).active.max_row-1,count)
    def test_18_dataset_excludes_demo_events(self):
        code,_,data=self.request('/api/demo/dataset.zip');self.assertEqual(code,200)
        with zipfile.ZipFile(io.BytesIO(data)) as z:self.assertEqual(json.loads(z.read('manifest.json'))['frames'],0);self.assertFalse(any(n.startswith('images/') for n in z.namelist()))
    def test_19_gif_animation_assets(self):
        from PIL import Image
        for p in ['/assets/animations/camera_06.gif','/assets/animations/camera_07.gif']:
            code,_,data=self.request(p);self.assertEqual(code,200);self.assertGreater(Image.open(io.BytesIO(data)).n_frames,1)
        self.assertEqual(self.request('/assets/photos/running_workplace_still.png')[0],200)
    def test_20_favicon_is_not_an_error(self):self.assertEqual(self.request('/favicon.ico')[0],204)
    def test_21_positive_model_duplicate_review_and_pseudo_labels(self):
        if not self.app.state()['model_ready']:self.fail('Model environment/checkpoint missing')
        encoded=base64.b64encode((ROOT/'demo/ff/assets/photos/trained_sample.jpg').read_bytes()).decode()
        code,_,raw=self.request('/api/demo/predict',{'image':encoded});self.assertEqual(code,200,raw.decode());result=json.loads(raw);self.assertEqual(result['prediction']['epochs'],26);self.assertFalse(result['prediction']['training_started']);self.assertIsNotNone(result['event']);id=result['event']['id']
        self.assertEqual(self.request('/api/demo/review',{'id':id,'status':'confirmed'})[0],200)
        again=json.loads(self.request('/api/demo/predict',{'image':encoded})[2]);self.assertTrue(again['duplicate']);self.assertEqual(again['event']['status'],'confirmed');self.assertEqual(len(self.get_state()['events']),6)
        self.request('/api/demo/reset',{});self.assertEqual(len(self.get_state()['events']),6);self.assertEqual(next(e for e in self.get_state()['events'] if e['id']==id)['status'],'confirmed')
        with zipfile.ZipFile(io.BytesIO(self.request('/api/demo/dataset.zip')[2])) as z:
            manifest=json.loads(z.read('manifest.json'));self.assertEqual(manifest['frames'],1);self.assertTrue(manifest['pseudo_labels']);self.assertIn('labels/'+id+'.txt',z.namelist())
            for line in z.read('labels/'+id+'.txt').decode().splitlines():self.assertEqual(len(line.split()),5);self.assertTrue(all(0<=float(x)<=1 for x in line.split()[1:]))
        (OUT/'model_result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    def test_22_negative_frame_and_busy_response(self):
        from PIL import Image
        image=Image.new('RGB',(320,320),(9,13,17));out=io.BytesIO();image.save(out,format='PNG');encoded=base64.b64encode(out.getvalue()).decode()
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            first=pool.submit(self.request,'/api/demo/predict',{'image':encoded});time.sleep(.3)
            second=self.request('/api/demo/predict',{'image':encoded});self.assertEqual(second[0],409)
            response=first.result();self.assertEqual(response[0],200,response[2].decode());self.assertIsNone(json.loads(response[2])['event']);self.assertEqual(len(self.get_state()['events']),5)
    def test_23_bad_base64_is_rejected(self):self.assertEqual(self.request('/api/demo/predict',{'image':'not:base64!!'})[0],400)
    def test_24_invalid_json_shape_and_non_image(self):
        self.assertEqual(self.request('/api/demo/review',raw=b'[]')[0],400)
        self.assertEqual(self.request('/api/demo/predict',{'image':base64.b64encode(b'not-an-image').decode()})[0],400)
    def test_25_parallel_pdf_exports(self):
        import pymupdf
        with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
            responses=list(pool.map(lambda _:self.request('/api/demo/report.pdf?download=1'),range(4)))
        for code,headers,data in responses:
            self.assertEqual(code,200);doc=pymupdf.open(stream=data,filetype='pdf');self.assertEqual(len(doc),1);self.assertIn('Отчёт по инцидентам',doc[0].get_text())
    def test_26_empty_journal_exports(self):
        import pymupdf
        from openpyxl import load_workbook
        with self.app.connect() as db:db.execute('DELETE FROM events')
        self.assertEqual(self.get_state()['pending'],0)
        doc=pymupdf.open(stream=self.request('/api/demo/report.pdf')[2],filetype='pdf');self.assertIn('0\n0\n0\n0\n0',doc[0].get_text())
        self.assertEqual(load_workbook(io.BytesIO(self.request('/api/demo/export.xlsx')[2])).active.max_row,1)
    def test_27_public_https_origin_and_host(self):
        self.app.args.public_origin='https://linq.user-link.xyz'
        try:
            self.assertEqual(self.request('/api/demo/state',headers={'Host':'linq.user-link.xyz'})[0],200)
            self.assertEqual(self.request('/api/demo/review',{'id':'EVT-2048','status':'confirmed'},headers={'Host':'linq.user-link.xyz','Origin':'https://linq.user-link.xyz'})[0],200)
            self.assertEqual(self.request('/api/demo/review',{'id':'EVT-2048','status':'rejected'},headers={'Host':'linq.user-link.xyz','Origin':'https://other.example'})[0],403)
            self.assertEqual(self.request('/api/demo/state',headers={'Host':'other.example'})[0],403)
        finally:del self.app.args.public_origin

if __name__=='__main__':
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(DemoTests);result=unittest.TextTestRunner(verbosity=2).run(suite)
    report={'tests_run':result.testsRun,'failures':[(str(t),s) for t,s in result.failures],'errors':[(str(t),s) for t,s in result.errors],'skipped':result.skipped,'passed':result.wasSuccessful(),'scope':'Isolated HTTP server, SQLite, assets, PDF, XLSX, actual 26-epoch CPU inference, duplicate prevention, dataset and errors.'}
    (OUT/'backend_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');sys.exit(0 if result.wasSuccessful() else 1)
